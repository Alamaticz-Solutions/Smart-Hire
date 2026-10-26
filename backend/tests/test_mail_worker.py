"""The mailbox poller reads and imports; replies are queued and sent at most once, only when switched on."""
from __future__ import annotations

import io
from email.message import EmailMessage

import docx
import pytest
from tests.conftest import FakeChat, login, payload

SETTINGS_BODY = {
    "email_enabled": 0, "imap_host": "", "imap_port": 993, "smtp_host": "", "smtp_port": 587,
    "email_user": "", "email_pass": "", "keywords": "resume,job", "drive_enabled": 0,
}


def _resume_email(sender: str = "candidate@example.com") -> bytes:
    document = docx.Document()
    document.add_paragraph("Asha Raman - Pega Senior System Architect")
    buffer = io.BytesIO()
    document.save(buffer)

    message = EmailMessage()
    message["From"] = f"Asha Raman <{sender}>"
    message["To"] = "careers@alamaticz.com"
    message["Subject"] = "Application - Pega architect"
    message["Message-ID"] = "<unit-test-1@example.com>"
    message["Date"] = "Tue, 10 Jun 2025 09:00:00 +0000"
    message.set_content("Please find my resume attached.")
    message.add_attachment(
        buffer.getvalue(), maintype="application",
        subtype="vnd.openxmlformats-officedocument.wordprocessingml.document", filename="asha_resume.docx",
    )
    return message.as_bytes()


class FakeImap:
    """Just enough IMAP for process_single_mailbox (one unseen message)."""

    raw = b""
    date_header = "Tue, 10 Jun 2025 09:00:00 +0000"
    sender = "candidate@example.com"

    def __init__(self, *args, **kwargs):
        pass

    def login(self, *args):
        return "OK", [b"logged in"]

    def select(self, *args):
        return "OK", [b"1"]

    def search(self, *args):
        return "OK", [b"1"]

    def fetch(self, num, query):
        if "HEADER" in query:
            headers = (
                f"Message-ID: <unit-test-1@example.com>\r\nSubject: Application\r\n"
                f"From: Asha <{FakeImap.sender}>\r\nDate: {FakeImap.date_header}\r\n\r\n"
            ).encode()
            return "OK", [(b"1 (BODY[HEADER])", headers), b")"]
        return "OK", [(b"1 (BODY[])", FakeImap.raw), b")"]

    def store(self, *args):
        return "OK", [b""]

    def logout(self):
        return "BYE", [b""]


class SmtpRecorder:
    sent: list = []

    def __init__(self, *args, **kwargs):
        pass

    def starttls(self):
        pass

    def login(self, *args):
        pass

    def sendmail(self, sender, recipient, body):
        SmtpRecorder.sent.append(recipient)

    def quit(self):
        pass


@pytest.fixture(autouse=True)
def _reset(monkeypatch):
    import app.services.email_worker as worker

    SmtpRecorder.sent = []
    FakeImap.date_header = "Tue, 10 Jun 2025 09:00:00 +0000"
    FakeImap.sender = "candidate@example.com"
    worker._guards.date_from = worker._guards.date_to = None
    worker._guards.skip_automated = True
    monkeypatch.setattr(worker.smtplib, "SMTP", SmtpRecorder)
    yield


def _run_worker(monkeypatch, sender="candidate@example.com"):
    import app.services.email_worker as worker

    FakeImap.raw = _resume_email(sender)
    FakeImap.sender = sender
    monkeypatch.setattr(worker.imaplib, "IMAP4_SSL", FakeImap)
    FakeChat.handler = staticmethod(lambda prompt: (
        '{"full_name": "Asha Raman", "total_experience": 8, "email": "asha@example.com", "skills": "Pega"}'
        if "expert resume and email parser" in prompt else "[]"
    ))
    worker.process_single_mailbox(
        "careers@alamaticz.com", "secret", "imap.example.com", 993, "smtp.example.com", 587,
        "resume,job", "professional", None, None, None,
    )
    return worker


def _candidates(client):
    return payload(client.get("/api/candidates", headers=login(client)))


def _replies(client):
    return payload(client.get("/api/integrations/replies", headers=login(client)))


def _smtp_settings(**extra):
    return {"email_user": "careers@alamaticz.com", "email_pass": "x", "smtp_host": "smtp.example.com",
            "smtp_port": 587, "imap_host": "imap.example.com", **extra}


def test_poller_imports_resume_and_only_queues_a_reply(client, monkeypatch):
    _run_worker(monkeypatch)

    rows = _candidates(client)
    assert [r["full_name"] for r in rows] == ["Asha Raman"]
    assert rows[0]["source"] == "uploaded from mail"
    assert SmtpRecorder.sent == []
    assert _replies(client)["counts"] == {"pending": 1, "sent": 0, "skipped": 0, "failed": 0}


def test_nothing_is_sent_while_auto_reply_is_off(client, monkeypatch):
    worker = _run_worker(monkeypatch)
    assert worker.send_pending_replies(_smtp_settings(auto_reply_enabled=0)) == 0
    assert SmtpRecorder.sent == []
    assert _replies(client)["counts"]["pending"] == 1


def test_queued_reply_is_sent_once_when_switched_on(client, monkeypatch):
    worker = _run_worker(monkeypatch)
    settings = _smtp_settings(auto_reply_enabled=1)

    assert worker.send_pending_replies(settings) == 1
    assert worker.send_pending_replies(settings) == 0  # second pass: nothing left
    assert len(SmtpRecorder.sent) == 1
    assert _replies(client)["counts"] == {"pending": 0, "sent": 1, "skipped": 0, "failed": 0}


def test_recipient_already_answered_is_not_emailed_again(client, monkeypatch):
    worker = _run_worker(monkeypatch)
    settings = _smtp_settings(auto_reply_enabled=1)
    worker.send_pending_replies(settings)

    candidate_id = _candidates(client)[0]["id"]
    recipient = _replies(client)["recent"][0]["recipient"]
    worker._queue_reply("another-message-uid", candidate_id, recipient, "Re: again", "careers@alamaticz.com")
    assert worker.send_pending_replies(settings) == 0
    assert len(SmtpRecorder.sent) == 1
    assert _replies(client)["counts"]["skipped"] == 1


def test_unsendable_reply_is_retried_then_given_up(client, monkeypatch):
    worker = _run_worker(monkeypatch)
    settings = {"auto_reply_enabled": 1}  # no credentials for this mailbox
    for _ in range(4):
        worker.send_pending_replies(settings)
    assert SmtpRecorder.sent == []
    assert _replies(client)["counts"]["failed"] == 1


def test_automated_senders_are_skipped(client, monkeypatch):
    _run_worker(monkeypatch, sender="invoice@vendor.com")
    assert _candidates(client) == []
    assert _replies(client)["counts"] == {"pending": 0, "sent": 0, "skipped": 0, "failed": 0}


def test_automated_sender_filter_can_be_turned_off(client, monkeypatch):
    import app.services.email_worker as worker

    worker._guards.skip_automated = False
    _run_worker(monkeypatch, sender="invoice@vendor.com")
    assert len(_candidates(client)) == 1


def test_mail_outside_the_date_window_is_left_alone(client, monkeypatch):
    import app.services.email_worker as worker

    worker._guards.date_from, worker._guards.date_to = "2025-07-01", "2025-07-31"
    _run_worker(monkeypatch)  # mail is dated 10 June
    assert _candidates(client) == []

    worker._guards.date_from, worker._guards.date_to = "2025-06-01", "2025-06-10"  # 'to' is inclusive
    _run_worker(monkeypatch)
    assert len(_candidates(client)) == 1  # not marked processed earlier, so it imports once in range


def test_message_is_not_processed_twice(client, monkeypatch):
    _run_worker(monkeypatch)
    _run_worker(monkeypatch)
    assert len(_candidates(client)) == 1
    assert _replies(client)["counts"]["pending"] == 1


def test_settings_validate_dates_and_default_auto_reply_off(client):
    admin = login(client)
    initial = payload(client.get("/api/integrations", headers=admin))
    assert initial["auto_reply_enabled"] == 0 and initial["skip_automated_senders"] == 1

    bad = client.post("/api/integrations", headers=admin, json={**SETTINGS_BODY, "mail_import_from": "10/06/2025"})
    assert bad.status_code == 400
    backwards = client.post("/api/integrations", headers=admin, json={
        **SETTINGS_BODY, "mail_import_from": "2025-06-10", "mail_import_to": "2025-06-01"})
    assert backwards.status_code == 400

    ok = client.post("/api/integrations", headers=admin, json={
        **SETTINGS_BODY, "mail_import_from": "2025-06-01", "mail_import_to": "2025-06-30", "auto_reply_enabled": 1})
    assert ok.status_code == 200
    saved = payload(client.get("/api/integrations", headers=admin))
    assert (saved["mail_import_from"], saved["mail_import_to"], saved["auto_reply_enabled"]) == ("2025-06-01", "2025-06-30", 1)


def test_reply_log_is_admin_only(client):
    assert client.get("/api/integrations/replies").status_code in (401, 403)


def test_failed_replies_are_listed_and_exported_to_excel(client, monkeypatch):
    import openpyxl

    worker = _run_worker(monkeypatch)
    for _ in range(3):  # no sending credentials -> gives up after the retry limit
        worker.send_pending_replies({"auto_reply_enabled": 1})

    log = _replies(client)
    assert [r["recipient"] for r in log["failed"]] == ["asha@example.com"]
    assert log["failed"][0]["error"]

    response = client.get("/api/integrations/replies/failed.xlsx", headers=login(client))
    assert response.status_code == 200
    rows = list(openpyxl.load_workbook(io.BytesIO(response.content)).active.values)
    assert rows[0][0] == "Candidate email" and rows[1][0] == "asha@example.com"


def test_candidate_keeps_the_date_its_mail_was_received(client, monkeypatch):
    _run_worker(monkeypatch)  # the fake mail is dated 10 Jun 2025 09:00 UTC
    row = _candidates(client)[0]
    assert str(row["mail_received_at"]).startswith("2025-06-10T09:00")
    assert any(c["col_key"] == "mail_received_at" for c in payload(client.get("/api/columns", headers=login(client)))["base"])
