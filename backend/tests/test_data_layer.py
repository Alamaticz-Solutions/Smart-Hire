"""Repository / service level tests for the ORM data layer."""
from __future__ import annotations

import io

import pytest
from sqlalchemy import select

from app.db.engine import session_scope
from app.models import Candidate, Job, User
from app.repositories import (
    CandidateRepository,
    JobCandidateRepository,
    JobRepository,
    UserRepository,
    invalidate_candidate_table_cache,
)
from app.services import columns as column_service
from tests.conftest import db_value, login, payload


# ── CandidateRepository ────────────────────────────────────────────────────
def test_candidate_insert_coerces_types_and_ignores_unknown_columns(client):
    with session_scope() as session:
        repo = CandidateRepository(session)
        candidate_id = repo.insert({
            "full_name": "Typed",
            "notice_period": 30,          # text column receives a number
            "total_experience": "4.5",    # float column receives text
            "availability_in_days": "15", # integer column receives text
            "email": "",                  # blank -> NULL
            "definitely_not_a_column": 1,
        })
        row = repo.get_row(candidate_id)
    assert row["notice_period"] == "30"
    assert row["total_experience"] == 4.5
    assert row["availability_in_days"] == 15
    assert row["email"] is None
    assert "definitely_not_a_column" not in row


def test_candidate_search_filters(client):
    with session_scope() as session:
        repo = CandidateRepository(session)
        repo.insert({"full_name": "Fast Joiner", "skills": "Pega CSSA", "pega_experience": 6,
                     "notice_period": "immediate", "current_location": "Chennai", "created_by": "u1"})
        repo.insert({"full_name": "Slow Joiner", "skills": "Pega", "pega_experience": 6,
                     "notice_period": "90", "current_location": "Pune", "created_by": "u1"})
        repo.insert({"full_name": "Other", "skills": "Java", "pega_experience": 0, "created_by": "u2"})

        names = lambda rows: sorted(r["full_name"] for r in rows)  # noqa: E731
        assert names(repo.search({"skills_keywords": ["pega"]}, None)) == ["Fast Joiner", "Slow Joiner"]
        assert names(repo.search({"skills_keywords": ["pega"], "notice_period_max": 30}, None)) == ["Fast Joiner"]
        assert names(repo.search({"current_location": "pune"}, None)) == ["Slow Joiner"]
        assert names(repo.search({"min_pega_exp": 5}, "u1")) == ["Fast Joiner", "Slow Joiner"]
        assert names(repo.search({"name_query": "other"}, "u1")) == []  # owner filter applied


def test_list_rows_pagination_and_count(client):
    with session_scope() as session:
        repo = CandidateRepository(session)
        for index in range(5):
            repo.insert({"full_name": f"C{index}", "created_by": "admin"})
        assert repo.count() == 5
        assert len(repo.list_rows(limit=2, offset=0)) == 2
        assert len(repo.list_rows(limit=10, offset=4)) == 1


def test_unmatched_excludes_linked_candidates(client):
    with session_scope() as session:
        candidates, jobs, links = CandidateRepository(session), JobRepository(session), JobCandidateRepository(session)
        linked = candidates.insert({"full_name": "Linked"})
        free = candidates.insert({"full_name": "Free"})
        job = jobs.add(Job(title="J", description="D"))
        links.add(job.id, linked, "reason")
        assert [r["full_name"] for r in candidates.list_unmatched_for_job(job.id)] == ["Free"]
        assert [r["id"] for r in candidates.list_for_job(job.id)] == [linked]
        assert free not in [r["id"] for r in candidates.list_for_job(job.id)]


# ── jobs ───────────────────────────────────────────────────────────────────
def test_job_aggregates_use_one_query_per_kind(client):
    with session_scope() as session:
        candidates, jobs, links = CandidateRepository(session), JobRepository(session), JobCandidateRepository(session)
        job = jobs.add(Job(title="J", description="D"))
        ids = [candidates.insert({"full_name": f"C{i}"}) for i in range(3)]
        links.add(job.id, ids[0], "r")
        links.add(job.id, ids[1], "r", status="selected")
        links.upsert_match(job.id, ids[2], "r")
        assert jobs.link_counts([job.id]) == {job.id: {"matched": 2, "selected": 1}}
        assert links.count_for_candidate(ids[0]) == 1

        links.clear_matched_for_job(job.id)  # selected links survive a re-match
        assert jobs.link_counts([job.id]) == {job.id: {"selected": 1}}


def test_user_lookup_is_case_insensitive(client):
    with session_scope() as session:
        users = UserRepository(session)
        assert users.get_by_username("ADMIN").username == "admin"
        assert users.username_exists("Sabari")
        assert not users.username_exists("nobody")


# ── custom columns service ─────────────────────────────────────────────────
def test_columns_service_roundtrip(client):
    with session_scope() as session:
        key = column_service.add_custom_column(session, "Notice Reason!", "Notice reason", "why leaving")
        assert key == "notice_reason"
        with pytest.raises(column_service.ColumnExistsError):
            column_service.add_custom_column(session, "notice_reason", "x", "")

    with session_scope() as session:
        repo = CandidateRepository(session)
        assert "notice_reason" in repo.column_names
        cid = repo.insert({"full_name": "Has Custom", "notice_reason": "relocation"})
        assert repo.get_row(cid)["notice_reason"] == "relocation"

    with session_scope() as session:
        column_service.remove_custom_column(session, "notice_reason")
        column_service.remove_custom_column(session, "full_name")  # built-ins are never dropped
        assert "notice_reason" not in CandidateRepository(session).column_names
        assert "full_name" in CandidateRepository(session).column_names


def test_clean_key():
    assert column_service.clean_key("Team Lead (EU)") == "team_lead_eu"


# ── bootstrap / seeding ────────────────────────────────────────────────────
def test_seed_data_present_and_admin_flags_repaired(client):
    from app.db.init_db import init_db

    with session_scope() as session:
        admin = UserRepository(session).get_by_username("admin")
        admin.is_admin = 0
        admin.is_approved = 0
    init_db()  # re-running restores the guaranteed flags
    assert db_value("SELECT is_admin FROM users WHERE username = 'admin'") == 1
    assert db_value("SELECT is_approved FROM users WHERE username = 'admin'") == 1
    assert db_value("SELECT COUNT(*) FROM team_members") == 4


def test_seeded_passwords_are_hashed_with_pbkdf2(client):
    assert db_value("SELECT password_hash FROM users WHERE username = 'admin'").startswith("pbkdf2_sha256$")


def test_production_seed_generates_unguessable_password(monkeypatch):
    import app.db.init_db as init_db_module

    monkeypatch.setattr(init_db_module, "DEBUG", False)
    monkeypatch.delenv("ADMIN_INITIAL_PASSWORD", raising=False)
    first = init_db_module._initial_password_hash("admin")
    from app.services.passwords import verify_password

    assert first.startswith("pbkdf2_sha256$")
    assert not verify_password("admin", first)

    monkeypatch.setenv("ADMIN_INITIAL_PASSWORD", "Correct-Horse-9")
    assert verify_password("Correct-Horse-9", init_db_module._initial_password_hash("admin"))


# ── vector store on SQLite ─────────────────────────────────────────────────
def test_vector_store_is_a_noop_without_postgres(client):
    from app.services import vector_store

    assert vector_store.is_enabled() is False
    assert vector_store.has_vectors() is False
    assert vector_store.index_documents([], "x.pdf") is False
    vector_store.delete_by_source("x.pdf")  # must not raise


# ── upload flow (LLM faked) ────────────────────────────────────────────────
def test_upload_rejects_unsupported_and_oversized_files(client):
    admin = login(client)
    bad = client.post("/api/upload", headers=admin, files={"file": ("virus.exe", b"MZ")})
    assert bad.status_code == 400

    huge = client.post("/api/upload", headers=admin, files={"file": ("big.pdf", b"x" * (26 * 1024 * 1024))})
    assert huge.status_code == 413


def test_resume_upload_creates_candidate_from_llm_extraction(client, monkeypatch):
    import docx

    from tests.conftest import FakeChat

    def respond(prompt: str) -> str:
        if "expert resume and email parser" in prompt:
            return ('{"full_name": "Meera Nair", "total_experience": 6, "pega_experience": 4, '
                    '"email": "meera@example.com", "phone": "+91 99999 11111", "skills": "Pega, Java", '
                    '"notice_period": "30", "source": "Resume Upload"}')
        if "expert resume formatter" in prompt:
            return '{"full_name": "Meera Nair", "job_title": "Pega Architect"}'
        return "[]"

    FakeChat.handler = staticmethod(respond)
    admin = login(client)

    document = docx.Document()
    document.add_paragraph("Meera Nair - Pega architect")
    buffer = io.BytesIO()
    document.save(buffer)

    response = client.post("/api/upload", headers=admin, files={"file": ("meera.docx", buffer.getvalue())})
    assert response.status_code == 200
    assert payload(response)["status"] == "processing"

    # TestClient runs background tasks before returning, so the row is final.
    rows = payload(client.get("/api/candidates", headers=admin))
    assert [r["full_name"] for r in rows] == ["Meera Nair"]
    assert rows[0]["total_experience"] == 6.0 and rows[0]["phone"] == "+919999911111"
    assert rows[0]["candidate_status"] == "New"


def test_failed_extraction_marks_placeholder_as_error_with_detail(client):
    import docx

    from tests.conftest import FakeChat

    FakeChat.handler = staticmethod(lambda prompt: "not json at all")
    admin = login(client)

    buffer = io.BytesIO()
    document = docx.Document()
    document.add_paragraph("whatever")
    document.save(buffer)
    client.post("/api/upload", headers=admin, files={"file": ("broken.docx", buffer.getvalue())})

    row = payload(client.get("/api/candidates", headers=admin))[0]
    assert row["candidate_status"] == "Error"
    assert row["error_detail"]
